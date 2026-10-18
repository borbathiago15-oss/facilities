"""O relatório do cliente em PDF (ReportLab, como o fechamento da Passagem) e
em Excel (openpyxl). Entra o dicionário de ``relatorios.cliente.dados`` — o
mesmo da tela — e as fotos escolhidas; sai o arquivo.

Régua do PDF: a primeira página responde sozinha (nota, variação e os quatro
indicadores); cor só para situação e sempre com a palavra junto; número que
não existe não vira zero ("sem movimento").
"""
import io

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from facilities import tempo

NAVY = colors.HexColor('#0f172a')
MARCA = colors.HexColor('#9d202c')
TINTA = colors.HexColor('#111827')
CINZA = colors.HexColor('#6b7280')
LINHA = colors.HexColor('#e5e7eb')
FUNDO = colors.HexColor('#f8fafc')
FAIXA = [colors.HexColor(c) for c in ('#9d202c', '#67a6dd', '#df8e24', '#a6bf38')]
COR_DA_FAIXA = {'boa': colors.HexColor('#15803d'), 'atencao': colors.HexColor('#b45309'),
                'critica': colors.HexColor('#b91c1c')}


def _esc(t):
    return (str(t if t is not None else '')
            .replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))


def _num(v, casas=1, sufixo=''):
    if v is None:
        return '—'
    texto = f'{v:.{casas}f}'.replace('.', ',') if isinstance(v, float) else str(v)
    return texto + sufixo


def _estilos():
    base = getSampleStyleSheet()

    def e(nome, **kw):
        return ParagraphStyle(nome, parent=base['Normal'], **kw)
    return {
        'marca': e('marca', fontSize=9, leading=12, textColor=MARCA, fontName='Helvetica-Bold'),
        'tit': e('tit', fontSize=22, leading=26, textColor=NAVY, fontName='Helvetica-Bold'),
        'sub': e('sub', fontSize=11, leading=15, textColor=CINZA),
        'secao': e('secao', fontSize=12.5, leading=16, textColor=NAVY,
                   fontName='Helvetica-Bold', spaceBefore=8, spaceAfter=4),
        'txt': e('txt', fontSize=9, leading=13, textColor=TINTA),
        'th': e('th', fontSize=7.5, leading=10, textColor=colors.white, fontName='Helvetica-Bold'),
        'cel': e('cel', fontSize=8, leading=11, textColor=TINTA),
        'pe': e('pe', fontSize=7, leading=9.5, textColor=CINZA),
        'kpi': e('kpi', fontSize=26, leading=30, textColor=NAVY, fontName='Helvetica-Bold',
                 alignment=TA_CENTER),
        'kpi_r': e('kpi_r', fontSize=8, leading=11, textColor=CINZA, alignment=TA_CENTER),
        'leg': e('leg', fontSize=7, leading=9, textColor=CINZA),
    }


def _tabela(cabecalho, linhas, larguras, est):
    dados = [[Paragraph(_esc(c), est['th']) for c in cabecalho]]
    dados += [[Paragraph(_esc(c), est['cel']) for c in linha] for linha in linhas]
    t = Table(dados, colWidths=larguras, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), NAVY), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LINEBELOW', (0, 0), (-1, -1), 0.4, LINHA),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, FUNDO]),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]))
    return t


def _faixa_da_marca(largura):
    t = Table([[''] * 4], colWidths=[largura / 4] * 4, rowHeights=[3])
    t.setStyle(TableStyle([('BACKGROUND', (i, 0), (i, 0), c) for i, c in enumerate(FAIXA)]
                          + [('LINEBELOW', (0, 0), (-1, -1), 0, colors.white)]))
    return t


def _imagem(dados, largura, altura):
    from reportlab.lib.utils import ImageReader
    leitor = ImageReader(io.BytesIO(dados))
    w, h = leitor.getSize()
    escala = min(largura / w, altura / h)
    return Image(io.BytesIO(dados), width=w * escala, height=h * escala)


def pdf(d, fotos, carregar_foto, recado=None):
    """``fotos``: a lista escolhida; ``carregar_foto(foto)`` → bytes (ou None)."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=16 * mm,
                            title=f'Relatório {d["local"].name} — {d["nome_do_mes"]}',
                            author='Grupo Trino · Facilities')
    est = _estilos()
    largura = A4[0] - 32 * mm
    h = [Paragraph('GRUPO TRINO · FACILITIES', est['marca']), Spacer(1, 2 * mm),
         _faixa_da_marca(largura), Spacer(1, 5 * mm),
         Paragraph(f'Relatório mensal — {_esc(d["local"].name)}', est['tit']),
         Paragraph(f'{_esc(d["cliente"] or "")}{" · " if d["cliente"] else ""}'
                   f'{_esc(d["nome_do_mes"])}', est['sub']), Spacer(1, 6 * mm)]

    nota = d['nota']
    faixa = nota['faixa']
    # A palavra vai sempre junto da cor (impressão em preto e branco). Sem
    # símbolo: a Helvetica do PDF não tem os círculos e imprime um quadrado.
    rotulo = d['faixas'].get(faixa, '') if faixa else 'sem movimento no mês'
    var = d['variacao']
    var_txt = ('igual ao mês anterior' if var == 0 else
               f'{"+" if var and var > 0 else ""}{_num(var)} ponto(s) sobre o mês anterior'
               if var is not None else 'sem comparação com o mês anterior')
    cartoes = [[Paragraph(_num(nota['nota']), est['kpi'])]
               + [Paragraph(_num(nota['indicadores'][k]['pct'], 1, '%'), est['kpi'])
                  for k in d['indicadores']],
               [Paragraph(f'Nota do local<br/>{_esc(rotulo)}', est['kpi_r'])]
               + [Paragraph(_esc(v), est['kpi_r']) for v in d['indicadores'].values()]]
    t = Table(cartoes, colWidths=[largura / 5] * 5)
    t.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), 0.6, LINHA),
                           ('INNERGRID', (0, 0), (-1, -1), 0.4, LINHA),
                           ('BACKGROUND', (0, 0), (0, -1), FUNDO),
                           ('TEXTCOLOR', (0, 0), (0, 0), COR_DA_FAIXA.get(faixa, NAVY)),
                           ('TOPPADDING', (0, 0), (-1, -1), 6),
                           ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
    h += [t, Spacer(1, 2 * mm), Paragraph(_esc(var_txt), est['pe']), Spacer(1, 4 * mm)]
    if recado:
        h += [Paragraph('Recado da coordenação', est['secao']), Paragraph(_esc(recado), est['txt'])]

    h.append(Paragraph('Rotinas por ambiente', est['secao']))
    if d['rotinas']:
        h.append(_tabela(['Ambiente', 'Previstas', 'No prazo', 'Com atraso', 'Não feitas',
                          'Aderência'],
                         [[r['ambiente'], r['previstas'], r['no_prazo'], r['com_atraso'],
                           r['nao_feitas'], _num(r['pct'], 1, '%')] for r in d['rotinas']],
                         [largura * x for x in (.36, .12, .12, .13, .13, .14)], est))
    else:
        h.append(Paragraph('Sem rotinas programadas no mês.', est['txt']))

    c = d['chamados']
    h += [Paragraph('Chamados', est['secao']),
          Paragraph(f'{c["abertos"]} aberto(s) no mês · {c["resolvidos"]} resolvido(s) · '
                    f'{_num(c["pct_no_prazo"], 1, "%")} no prazo · tempo médio de solução '
                    f'{_num(c["tempo_medio_h"], 1, " h")} · {c["pendentes_no_fim"]} em aberto no '
                    'fim do mês.', est['txt']), Spacer(1, 2 * mm)]
    if c['lista']:
        h.append(_tabela(['Nº', 'Assunto', 'Tipo', 'Aberto', 'Resolvido', 'Situação'],
                         [[x['numero'], x['titulo'], x['categoria'],
                           tempo.fmt(x['aberto_em'], '%d/%m %H:%M'),
                           tempo.fmt(x['resolvido_em'], '%d/%m %H:%M'),
                           x['situacao'] + (' · no prazo' if x['no_prazo'] else
                                            ' · fora do prazo' if x['no_prazo'] is False else '')]
                          for x in c['lista']],
                         [largura * x for x in (.1, .34, .14, .13, .13, .16)], est))
    n = d['ncs']
    h += [Paragraph('O que não estava conforme', est['secao']),
          Paragraph(f'{n["total"]} item(ns) não conforme(s) nos checklists; '
                    f'{n["corrigidas"]} já corrigido(s) por chamado.', est['txt'])]
    if n['itens']:
        h.append(_tabela(['Item mais frequente', 'Vezes'], [[i, q] for i, q in n['itens']],
                         [largura * .8, largura * .2], est))
    o, v, cb = d['ocorrencias'], d['visitas'], d['cobrancas']
    h += [Paragraph('Avisos do público, visitas e cobranças', est['secao']),
          _tabela(['O quê', 'No mês'], [
              ['Avisos do público pelo QR', f'{o["total"]} recebido(s), {o["resolvidas"]} '
               f'resolvido(s), {_num(o["pct_no_prazo"], 1, "%")} no prazo, '
               f'tempo médio {_num(o["tempo_medio_min"], 0, " min")}'],
              ['Visitas de supervisão', f'{v["feitas"]} de {v["programadas"]} feitas, '
               f'{v["conformes"]} sem nenhum item não conforme'],
              ['Suas cobranças', f'{cb["abertas"]} aberta(s), {cb["respondidas"]} respondida(s) '
               f'({cb["no_prazo"]} no prazo), {cb["encerradas"]} encerrada(s)']],
              [largura * .35, largura * .65], est)]

    if fotos:
        h.append(Paragraph('Fotos do mês', est['secao']))
        celulas, linha = [], []
        lado = (largura - 6 * mm) / 3
        for f in fotos:
            dados_foto = carregar_foto(f)
            if not dados_foto:
                continue
            try:
                img = _imagem(dados_foto, lado - 4, lado * .6)
            except Exception:
                continue
            linha.append([img, Paragraph(_esc(f['legenda']), est['leg'])])
            if len(linha) == 3:
                celulas.append(linha)
                linha = []
        if linha:
            celulas.append(linha + [''] * (3 - len(linha)))
        if celulas:
            t = Table(celulas, colWidths=[lado + 2 * mm] * 3)
            t.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                   ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
            h.append(KeepTogether([t]) if len(celulas) == 1 else t)

    h += [Spacer(1, 3 * mm),
          Paragraph(f'Gerado em {tempo.fmt(d["gerado_em"])} pelo Trino Facilities. Os números '
                    'saem dos registros do sistema (checklists lidos no ponto pelo QR, chamados e '
                    'avisos), os mesmos do portal do cliente.', est['pe'])]

    def rodape(canvas, _doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(CINZA)
        canvas.drawString(16 * mm, 9 * mm, f'{d["local"].name} · {d["nome_do_mes"]}')
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, f'página {canvas.getPageNumber()}')
        canvas.restoreState()
    doc.build(h, onFirstPage=rodape, onLaterPages=rodape)
    return buf.getvalue()


def excel(d):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    cab = Font(bold=True, color='FFFFFF')
    fundo = PatternFill('solid', fgColor='0F172A')

    def aba(ws, cabecalho, linhas):
        ws.append(cabecalho)
        for c in ws[1]:
            c.font, c.fill = cab, fundo
        for linha in linhas:
            ws.append(list(linha))
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                12, min(60, max(len(str(c.value or '')) for c in col) + 2))

    ws = wb.active
    ws.title = 'Resumo'
    nota = d['nota']
    aba(ws, ['Indicador', 'Valor', 'Base'], [
        ['Local', d['local'].name, ''], ['Cliente', d['cliente'] or '', ''],
        ['Mês', d['nome_do_mes'], ''],
        ['Nota do local', nota['nota'], d['faixas'].get(nota['faixa'], 'sem movimento')],
        ['Variação sobre o mês anterior', d['variacao'], ''],
    ] + [[rot, nota['indicadores'][k]['pct'],
          f'{nota["indicadores"][k]["bom"]} de {nota["indicadores"][k]["base"]}']
         for k, rot in d['indicadores'].items()])
    aba(wb.create_sheet('Rotinas por ambiente'),
        ['Ambiente', 'Previstas', 'No prazo', 'Com atraso', 'Não feitas', 'Aderência (%)'],
        [[r['ambiente'], r['previstas'], r['no_prazo'], r['com_atraso'], r['nao_feitas'],
          r['pct']] for r in d['rotinas']])
    aba(wb.create_sheet('Chamados'),
        ['Nº', 'Assunto', 'Tipo', 'Aberto em', 'Prazo', 'Resolvido em', 'Situação', 'No prazo'],
        [[x['numero'], x['titulo'], x['categoria'], tempo.fmt(x['aberto_em'], vazio=''),
          tempo.fmt(x['prazo'], vazio=''), tempo.fmt(x['resolvido_em'], vazio=''),
          x['situacao'], {True: 'sim', False: 'não'}.get(x['no_prazo'], '')]
         for x in d['chamados']['lista']])
    aba(wb.create_sheet('Não conformidades'), ['Item', 'Vezes'], d['ncs']['itens'])
    o, v, cb = d['ocorrencias'], d['visitas'], d['cobrancas']
    aba(wb.create_sheet('Avisos, visitas, cobranças'), ['O quê', 'Valor'], [
        ['Avisos do público recebidos', o['total']], ['Avisos resolvidos', o['resolvidas']],
        ['Avisos no prazo (%)', o['pct_no_prazo']],
        ['Tempo médio de atendimento (min)', o['tempo_medio_min']],
        ['Visitas programadas', v['programadas']], ['Visitas feitas', v['feitas']],
        ['Visitas sem não conformidade', v['conformes']],
        ['Cobranças abertas', cb['abertas']], ['Cobranças respondidas', cb['respondidas']],
        ['Respondidas no prazo', cb['no_prazo']], ['Cobranças encerradas', cb['encerradas']]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
