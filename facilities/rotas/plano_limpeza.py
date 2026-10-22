"""O plano de limpeza do cliente (planilha "PLANO DE LIMPEZA E CONSERVAÇÃO")
vira blocos, checklists e rotas do local (decisões do dono, 22/10/2026).

Três passos, cada um sem o anterior saber do seguinte:

1. ``ler(arquivo)`` → ``PlanoLido``: lê a planilha como ela é (uma aba por
   rota: cabeçalho com responsável, entrada, intervalo, largada e os dias;
   depois uma linha por atividade com material, EPI, frequência e, por dia,
   "HH:MMh / HH:MMh" ou "x" quando não há). Puro: nada de banco. O que não
   entende vai para ``avisos`` em vez de parar.
2. ``planejar(plano, local)`` → ``Previa``: decide o que seria criado — os
   BLOCOS (um ponto com QR por bloco, por decisão do dono; desmembrar vem
   depois), os checklists (um por atividade distinta) e as rotas com os
   itens — sem gravar. É o que a tela e o ``--dry-run`` mostram.
3. ``aplicar(previa, local, autor)`` grava, reaproveitando setor, bloco,
   ponto e checklist que já existem (como a rota pronta faz). Rota com o
   mesmo nome já no local: recusa, a não ser ``substituir`` (a antiga é
   desativada e renomeada; o histórico fica).

A escala (quem faz a rota) NÃO é criada: nome de pessoa não entra pelo
plano. A prévia só lembra "Responsável: FULANO" para a gestão escalar.

Regras do bloco (de onde sai o ponto/QR de cada atividade), nesta ordem:

| A atividade...                                   | Setor         | Bloco (ambiente)                 |
|--------------------------------------------------|---------------|----------------------------------|
| começa com "Continuação das atividades..."       | o da atividade de referência: a última anterior que não é de bloco fixo (gaiolas, lixo, utensílios); senão, a imediatamente anterior |
| "PREFIXO EM MAIÚSCULAS - Sub-bloco: o que fazer" | PREFIXO       | Sub-bloco ("Andar administrativo", "Mezanino", "Térreo"; "Andar térreo" vira "Térreo") |
| "PREFIXO EM MAIÚSCULAS - o que fazer"            | PREFIXO       | o próprio prefixo (ex.: "Pátio externo", "Lactalis") |
| "Limpeza e guarda dos utensílios..."             | Operação      | DML (guarda de utensílios)       |
| "Realizar degelo e troca de EPI's..."            | Operação      | DML (guarda de utensílios)       |
| "Recolhe o lixo não reciclável para a Central..."| Operação      | Central de resíduos              |
| "Recolhe/Disponibilizar/Retornar as gaiolas..."  | Operação      | Doca 09 / triagem de resíduos    |
| "Limpeza das antecâmaras D1, câmaras 1A..."      | Operação      | Câmaras D1                       |
| "Limpeza da antecâmara D2, câmaras 5A..."        | Operação      | Câmaras D2                       |
| "...garrafas de café das salas administrativas"  | GALPÃO FRIO   | Andar administrativo             |
| nenhuma das anteriores                           | Operação      | Operação (geral), com aviso      |

Checklist por atividade distinta (o texto normalizado — sem acento, sem
maiúsculas, sem espaços repetidos — é a chave): nome = o texto sem o prefixo
do bloco, até 150 letras; descrição = o texto completo; materiais e EPI da
planilha. Perguntas: quando há dois-pontos, a parte depois dele quebrada em
", ", " + " e " / " (ex.: "Limpeza de WCs, corredor e salas, copa e degelo"
→ WCs · corredor e salas · copa e degelo); senão, uma só: "Feito conforme o
plano". Todas sim/não, obrigatórias, problema quando "não".

Item da rota por atividade: ponto do bloco, checklist, horário = início,
tempo previsto = fim − início (atravessando a meia-noite quando precisa),
``dias_semana`` quando algum dia é "x", ``horarios_por_dia`` quando um dia
começa em horário diferente dos outros (o sábado, em geral).
"""
import io
import re
import unicodedata
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, time

from facilities.extensions import db
from facilities.rotas import servico
from facilities.rotas.models import ChecklistModelo, Rota
from facilities.rotas.servico import DIAS_DA_SEMANA, RegraDaRota

PLANILHA_MAX = 5 * 1024 * 1024
TOLERANCIA_DO_PLANO = 15
CHAVE_DA_PLANILHA = re.compile(r'^planos/(\d+)/[0-9a-f]{32}\.xlsx$')
# Onde está a planilha da Trino Frio empacotada (``semear_trino_frio``).
PLANILHA_TRINO_FRIO = 'data/plano_trino_frio.xlsx'

_DIA_POR_NOME = {'SEG': '1', 'TER': '2', 'QUA': '3', 'QUI': '4', 'SEX': '5', 'SAB': '6',
                 'SÁB': '6', 'DOM': '7'}
_ORDEM_DA_SEMANA = ['SEGUNDA', 'TERCA', 'QUARTA', 'QUINTA', 'SEXTA', 'SABADO', 'DOMINGO']
_HORA = re.compile(r'(\d{1,2})([:.]?)(\d{2})\s*[hH]?')
_PREFIXO = re.compile(r'^([A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-ZÁÉÍÓÚÂÊÔÃÕÇ\s]{1,30}?)\s+-\s+(.+)$', re.S)
_SEPARADOR_DE_ITENS = re.compile(r'\s*/\s*|,\s+|\s\+\s')
_VERBO_INICIAL = re.compile(r'^(repasse\s+d[aeo]s?\s+limpeza\s+d[aeo]s?|repasse\s+d[aeo]s?|'
                            r'limpeza\s+d[aeo]s?|abastecimento\s+d[aeo]s?)\s+', re.I)
_SUBBLOCO = {'andar terreo': 'Térreo', 'terreo': 'Térreo',
             'andar administrativo': 'Andar administrativo', 'mezanino': 'Mezanino'}
SETOR_OPERACAO = 'Operação'
# (pistas no texto normalizado, setor, bloco). A primeira que casa vale.
BLOCOS_FIXOS = (
    (('guarda dos utensilios',), SETOR_OPERACAO, 'DML (guarda de utensílios)'),
    (('degelo e troca de epi',), SETOR_OPERACAO, 'DML (guarda de utensílios)'),
    (('lixo nao reciclavel', 'central de residuos'), SETOR_OPERACAO, 'Central de resíduos'),
    (('recolhe as gaiolas', 'disponibilizar gaiolas', 'retornar todas as gaiolas',
      'triagem dos residuos', 'triagem de reciclaveis'), SETOR_OPERACAO,
     'Doca 09 / triagem de resíduos'),
    (('antecamaras d1', 'antecamara d1', 'camaras 1a'), SETOR_OPERACAO, 'Câmaras D1'),
    (('antecamara d2', 'camaras 5a'), SETOR_OPERACAO, 'Câmaras D2'),
    (('salas administrativas',), 'GALPÃO FRIO', 'Andar administrativo'),
)
BLOCOS_DE_APOIO = ('DML (guarda de utensílios)', 'Central de resíduos',
                   'Doca 09 / triagem de resíduos')
BLOCO_SEM_LUGAR = 'Operação (geral)'
ITEM_UNICO = 'Feito conforme o plano'


# ── 1. Leitura ────────────────────────────────────────────────────────────

@dataclass
class AtividadeLida:
    descricao: str
    material: str = ''
    epi: str = ''
    frequencia: str = ''
    turno_bloco: str = ''
    linha: int = 0
    # dia ('1'..'7') → (início, fim). Dia ausente = "x" na planilha.
    horarios_por_dia: dict = field(default_factory=dict)


@dataclass
class RotaLida:
    numero: int
    responsavel: str = ''
    entrada: time = None
    largada: time = None
    intervalo: tuple = None            # (início, fim) ou None
    dias: str = ''                     # vocabulário de DIAS_DA_SEMANA: '123456'
    observacoes: str = ''
    atividades: list = field(default_factory=list)


@dataclass
class PlanoLido:
    rotas: list = field(default_factory=list)
    avisos: list = field(default_factory=list)
    revisoes: int = 0


def _normal(texto):
    s = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()


def _texto(valor):
    return re.sub(r'[ \t]+', ' ', str(valor).strip()) if valor is not None else ''


def _horas(texto, onde, avisos):
    """Os horários de uma célula ("06:00h\n06:30h", "1600h") → [time]."""
    saida = []
    for h, sep, m, in _HORA.findall(str(texto or '')):
        h, m = int(h), int(m)
        if not (0 <= h < 24 and 0 <= m < 60):
            avisos.append(f'{onde}: "{texto}" não é um horário válido.')
            continue
        if not sep:
            avisos.append(f'{onde}: "{h}{m:02d}h" entendido como {h:02d}:{m:02d}.')
        saida.append(time(h, m))
    return saida


def _dias_do_texto(texto):
    """"SEGUNDA A SÁBADO" → '123456'; "DOMINGO A SEXTA" → '123457'."""
    nomes = [n for n in re.split(r'\s+A\s+', _normal(texto).upper()) if n]
    if len(nomes) != 2 or nomes[0] not in _ORDEM_DA_SEMANA or nomes[1] not in _ORDEM_DA_SEMANA:
        return ''
    i, f = _ORDEM_DA_SEMANA.index(nomes[0]), _ORDEM_DA_SEMANA.index(nomes[1])
    dias = []
    while True:
        dias.append(str(i + 1))
        if i == f:
            break
        i = (i + 1) % 7
    return ''.join(sorted(dias))


def _ler_aba(ws, numero, plano):
    rota = RotaLida(numero=numero)
    nome = f'Rota {numero}'
    avisos = plano.avisos
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    colunas_dos_dias, dias_do_texto, turno_bloco = {}, '', ''
    i = 0
    while i < len(linhas):
        linha = linhas[i]
        textos = [_texto(c) for c in linha]
        maiusc = [t.upper() for t in textos]
        if 'RESPONSÁVEL' in maiusc and i + 1 < len(linhas):
            valores = [_texto(c) for c in linhas[i + 1]]
            for col, titulo in enumerate(maiusc):
                v = valores[col] if col < len(valores) else ''
                if titulo == 'RESPONSÁVEL':
                    rota.responsavel = v
                elif titulo == 'ENTRADA':
                    hs = _horas(v, f'{nome}, entrada', avisos)
                    rota.entrada = hs[0] if hs else None
                elif titulo == 'LARGADA':
                    hs = _horas(v, f'{nome}, largada', avisos)
                    rota.largada = hs[0] if hs else None
                elif titulo == 'INTERVALO':
                    hs = _horas(v, f'{nome}, intervalo', avisos)
                    rota.intervalo = (hs[0], hs[1]) if len(hs) >= 2 else None
                elif ' A ' in f' {titulo} ' and _dias_do_texto(titulo):
                    dias_do_texto = _dias_do_texto(titulo)
            i += 2
            continue
        if maiusc and maiusc[0] == 'ATIVIDADE':
            colunas_dos_dias = {col: _DIA_POR_NOME[t] for col, t in enumerate(maiusc)
                                if t in _DIA_POR_NOME}
            rota.dias = ''.join(sorted(colunas_dos_dias.values()))
            # "1° TURNO" começa na própria linha de cabeçalho (célula mesclada).
            turno_bloco = textos[4] if len(textos) > 4 and textos[4] else turno_bloco
            if dias_do_texto and dias_do_texto != rota.dias:
                avisos.append(f'{nome}: o cabeçalho diz um período e as colunas trazem outro; '
                              'valem as colunas.')
            i += 1
            continue
        primeiro = textos[0] if textos else ''
        if not primeiro or not colunas_dos_dias:
            i += 1
            continue
        if primeiro.upper().startswith('INTERVALO'):
            hs = _horas(primeiro, f'{nome}, linha do intervalo', avisos)
            if len(hs) >= 2:
                if rota.intervalo and rota.intervalo != (hs[0], hs[1]):
                    avisos.append(f'{nome}: o intervalo do cabeçalho '
                                  f'({rota.intervalo[0].strftime("%H:%M")} às '
                                  f'{rota.intervalo[1].strftime("%H:%M")}) difere do da grade '
                                  f'({hs[0].strftime("%H:%M")} às {hs[1].strftime("%H:%M")}); '
                                  'vale o da grade, que é o que as atividades seguem.')
                rota.intervalo = (hs[0], hs[1])
            i += 1
            continue
        if primeiro.upper().startswith('OBSERVAÇ'):
            rota.observacoes = '\n'.join(t for t in textos[1:] if t).strip()
            i += 1
            continue
        atividade = AtividadeLida(descricao=primeiro.strip(' .\n'), linha=i + 1,
                                  material=textos[1] if len(textos) > 1 else '',
                                  epi=textos[2] if len(textos) > 2 else '',
                                  frequencia=textos[3] if len(textos) > 3 else '')
        # "1° TURNO" / "2° TURNO" está só na primeira linha do grupo (célula
        # mesclada): vale até aparecer o próximo.
        turno_bloco = (textos[4] if len(textos) > 4 and textos[4] else turno_bloco)
        atividade.turno_bloco = turno_bloco
        for col, dia in colunas_dos_dias.items():
            v = textos[col] if col < len(textos) else ''
            if not v or v.lower().strip('. ') in ('x', 'xx'):
                continue
            onde = f'{nome}, linha {i + 1}, {DIAS_DA_SEMANA[dia]}'
            hs = _horas(v, onde, avisos)
            if len(hs) != 2:
                avisos.append(f'{onde}: "{v}" não tem início e fim; dia ignorado.')
                continue
            atividade.horarios_por_dia[dia] = (hs[0], hs[1])
        if atividade.horarios_por_dia:
            rota.atividades.append(atividade)
        else:
            avisos.append(f'{nome}, linha {i + 1}: "{primeiro[:60]}" sem horário em nenhum '
                          'dia; ignorada.')
        i += 1
    for campo in ('material', 'epi'):
        for a in rota.atividades:
            if _normal(getattr(a, campo)) in ('x', 'xx', '-'):
                setattr(a, campo, '')
    if rota.entrada is None or rota.largada is None:
        avisos.append(f'{nome}: sem entrada ou largada; a rota não entra.')
        return None
    if not rota.atividades:
        avisos.append(f'{nome}: nenhuma atividade com horário; a rota não entra.')
        return None
    return rota


def ler(arquivo):
    """A planilha (caminho, bytes ou arquivo aberto) → ``PlanoLido``.

    Puro: não toca o banco. Abas "ROTA N" viram rotas; "REVISÕES" só é
    contada; "ROTAs MODELO" (em branco) é ignorada."""
    import openpyxl
    if isinstance(arquivo, (bytes, bytearray)):
        arquivo = io.BytesIO(arquivo)
    try:
        wb = openpyxl.load_workbook(arquivo, data_only=True, read_only=True)
    except Exception as erro:
        raise RegraDaRota(f'Não consegui abrir a planilha ({type(erro).__name__}). Envie o '
                          'arquivo .xlsx do plano de limpeza.')
    plano = PlanoLido()
    try:
        for ws in wb.worksheets:
            titulo = _texto(ws.title).upper()
            m = re.fullmatch(r'ROTA\s*(\d+)', titulo)
            if m and int(m.group(1)) > 0:
                rota = _ler_aba(ws, int(m.group(1)), plano)
                if rota is not None:
                    plano.rotas.append(rota)
            elif titulo.startswith('REVIS'):
                plano.revisoes = sum(1 for r in ws.iter_rows(values_only=True)
                                     if any(isinstance(c, datetime) for c in r))
            elif 'MODELO' not in titulo:
                plano.avisos.append(f'Aba "{ws.title}" ignorada: não é uma rota.')
    finally:
        wb.close()
    plano.rotas.sort(key=lambda r: r.numero)
    if not plano.rotas:
        raise RegraDaRota('A planilha não tem nenhuma aba "ROTA N" com atividades.')
    return plano


# ── 2. Planejamento ───────────────────────────────────────────────────────

@dataclass
class BlocoPrevisto:
    setor: str
    nome: str                           # o ambiente (um ponto com QR por bloco)
    atividades: int = 0
    existe: bool = False

    @property
    def chave(self):
        return (_normal(self.setor), _normal(self.nome))

    @property
    def rotulo(self):
        return self.nome if _normal(self.nome) == _normal(self.setor) \
            else f'{self.setor} — {self.nome}'


@dataclass
class ChecklistPrevisto:
    nome: str
    descricao: str
    materiais: str = ''
    epi: str = ''
    itens: list = field(default_factory=list)
    existe: bool = False


@dataclass
class ItemPrevisto:
    ordem: int
    descricao: str
    bloco: BlocoPrevisto
    checklist: ChecklistPrevisto
    horario: str
    tempo_previsto_min: int = None
    dias_semana: str = None
    horarios_por_dia: dict = field(default_factory=dict)
    turno_bloco: str = ''


@dataclass
class RotaPrevista:
    numero: int
    nome: str
    turno: str
    hora_inicio: str
    hora_fim: str
    dias_semana: str
    responsavel: str = ''
    intervalo_inicio: str = ''
    intervalo_fim: str = ''
    observacoes: str = ''
    tolerancia_min: int = TOLERANCIA_DO_PLANO
    itens: list = field(default_factory=list)
    existe: bool = False


@dataclass
class Previa:
    blocos: list = field(default_factory=list)
    checklists: list = field(default_factory=list)
    rotas: list = field(default_factory=list)
    avisos: list = field(default_factory=list)

    @property
    def responsaveis(self):
        return [(r.numero, r.responsavel) for r in self.rotas if r.responsavel]


def _bloco_da_atividade(descricao):
    """→ (setor, bloco, texto sem o prefixo) ou None quando não há pista."""
    texto = descricao.strip(' .')
    m = _PREFIXO.match(texto)
    if m and m.group(1).strip() == m.group(1).strip().upper():
        prefixo, resto = m.group(1).strip(), m.group(2).strip()
        sub, sep, _acao = resto.partition(':')
        if sep and len(sub.split()) <= 4:
            nome = _SUBBLOCO.get(_normal(sub), sub.strip())
            return prefixo, nome, resto
        return prefixo, prefixo.title(), resto
    n = _normal(texto)
    for pistas, setor, bloco in BLOCOS_FIXOS:
        if any(p in n for p in pistas):
            return setor, bloco, texto
    return None


def _nome_do_checklist(texto):
    nome = re.sub(r'\s+', ' ', texto.split('\n', 1)[0]).strip(' .')
    if len(nome) > 150:
        nome = nome[:149].rsplit(' ', 1)[0] + '…'
    return nome


def _itens_do_checklist(texto):
    """As perguntas: a parte depois dos dois-pontos quebrada em partes, ou
    uma só ("Feito conforme o plano")."""
    primeira_linha = texto.split('\n', 1)[0]
    _antes, sep, depois = primeira_linha.partition(':')
    partes = []
    if sep:
        for n, p in enumerate(_SEPARADOR_DE_ITENS.split(depois)):
            p = p.strip(' .;')
            if n == 0:
                p = _VERBO_INICIAL.sub('', p)
            if p:
                partes.append(p[0].upper() + p[1:])
    if not partes:
        partes = [ITEM_UNICO]
    itens = []
    for p in partes:
        p = p[:300]
        itens.append({'texto': p, 'tipo': 'sim_nao', 'obrigatorio': True, 'problema_se': 'nao',
                      'problema': (f'{p} com problema' if p != ITEM_UNICO
                                   else 'Não feito conforme o plano')[:120]})
    return servico.normalizar_itens(itens)


def _minutos(ini, fim):
    return ((fim.hour * 60 + fim.minute) - (ini.hour * 60 + ini.minute)) % (24 * 60)


def _dentro_do_turno(h, entrada, largada):
    a, b, x = entrada.hour * 60 + entrada.minute, largada.hour * 60 + largada.minute, \
        h.hour * 60 + h.minute
    if b <= a:
        b += 24 * 60
        if x < a:
            x += 24 * 60
    return a <= x <= b


def _dentro_do_intervalo(h, intervalo, entrada):
    if not intervalo:
        return False
    a = entrada.hour * 60 + entrada.minute
    ii, ff, x = [(t.hour * 60 + t.minute) for t in (*intervalo, h)]
    ii, ff, x = [v + 24 * 60 if v < a else v for v in (ii, ff, x)]
    if ff <= ii:
        ff += 24 * 60
    return ii <= x < ff


def planejar(plano, local=None):
    """O que seria criado no local: blocos, checklists e rotas. Não grava.
    Com ``local``, marca o que já existe lá (setor/bloco, checklist, rota)."""
    previa = Previa(avisos=list(plano.avisos))
    blocos, checklists = {}, {}

    def bloco(setor, nome):
        chave = (_normal(setor), _normal(nome))
        if chave not in blocos:
            blocos[chave] = BlocoPrevisto(setor=setor, nome=nome)
        return blocos[chave]

    def checklist(texto, atividade):
        chave = _normal(texto)
        if chave not in checklists:
            nome = _nome_do_checklist(texto)
            usados = {_normal(c.nome) for c in checklists.values()}
            base, n = nome, 2
            while _normal(nome) in usados:
                nome = f'{base[:145]} ({n})'
                n += 1
            checklists[chave] = ChecklistPrevisto(
                nome=nome, descricao=atividade.descricao.strip(),
                materiais=atividade.material.strip()[:servico.MATERIAIS_MAX],
                epi=atividade.epi.strip()[:servico.MATERIAIS_MAX],
                itens=_itens_do_checklist(texto))
        return checklists[chave]

    for r in plano.rotas:
        nome_rota = f'Rota {r.numero}'
        rota = RotaPrevista(
            numero=r.numero, nome=nome_rota, responsavel=r.responsavel,
            turno=('Noturno' if r.largada <= r.entrada
                   else 'Tarde' if r.entrada >= time(12, 0) else 'Diurno'),
            hora_inicio=r.entrada.strftime('%H:%M'), hora_fim=r.largada.strftime('%H:%M'),
            dias_semana=r.dias or '1234567',
            intervalo_inicio=r.intervalo[0].strftime('%H:%M') if r.intervalo else '',
            intervalo_fim=r.intervalo[1].strftime('%H:%M') if r.intervalo else '',
            observacoes=r.observacoes)
        referencia = None                   # (bloco, checklist) para "continuação"
        for a in r.atividades:
            onde = f'{nome_rota}, linha {a.linha}'
            if _normal(a.descricao).startswith('continuacao'):
                if referencia is None:
                    previa.avisos.append(f'{onde}: "Continuação" sem atividade anterior; '
                                         'ignorada.')
                    continue
                b, ck = referencia
            else:
                achado = _bloco_da_atividade(a.descricao)
                if achado is None:
                    previa.avisos.append(f'{onde}: não achei o lugar de "{a.descricao[:50]}"; '
                                         f'foi para "{BLOCO_SEM_LUGAR}". Ajuste na rota.')
                    setor, nome_bloco, texto = SETOR_OPERACAO, BLOCO_SEM_LUGAR, a.descricao
                else:
                    setor, nome_bloco, texto = achado
                b, ck = bloco(setor, nome_bloco), checklist(texto, a)
            if b.nome not in BLOCOS_DE_APOIO or referencia is None:
                referencia = (b, ck)
            # O horário comum (o mais frequente entre os dias) e os dias que diferem.
            pares = Counter(a.horarios_por_dia.values())
            comum = max(pares, key=lambda p: (pares[p], -list(a.horarios_por_dia.values())
                                               .index(p)))
            ini, fim = comum
            if not _dentro_do_turno(ini, r.entrada, r.largada):
                previa.avisos.append(f'{onde}: {ini.strftime("%H:%M")} fica fora do turno '
                                     f'({rota.hora_inicio} às {rota.hora_fim}); ignorada.')
                continue
            if _dentro_do_intervalo(ini, r.intervalo, r.entrada):
                previa.avisos.append(f'{onde}: {ini.strftime("%H:%M")} cai no intervalo; o '
                                     'agendador não vai gerar essa tarefa.')
            minutos = _minutos(ini, fim)
            if minutos == 0:
                previa.avisos.append(f'{onde}: início e fim iguais; sem tempo previsto.')
                minutos = None
            elif minutos > servico.TEMPO_PREVISTO_MAX:
                previa.avisos.append(f'{onde}: {minutos} min passa do máximo; ficou '
                                     f'{servico.TEMPO_PREVISTO_MAX}.')
                minutos = servico.TEMPO_PREVISTO_MAX
            por_dia = {}
            for d, (i2, _f2) in sorted(a.horarios_por_dia.items()):
                if i2 == ini:
                    continue
                if not _dentro_do_turno(i2, r.entrada, r.largada):
                    previa.avisos.append(
                        f'{onde}, {DIAS_DA_SEMANA[d]}: {i2.strftime("%H:%M")} fica fora do '
                        f'turno; nesse dia vale {ini.strftime("%H:%M")}.')
                    continue
                por_dia[d] = i2.strftime('%H:%M')
            dias = ''.join(sorted(a.horarios_por_dia))
            item = ItemPrevisto(
                ordem=len(rota.itens) + 1, descricao=a.descricao, bloco=b, checklist=ck,
                horario=ini.strftime('%H:%M'), tempo_previsto_min=minutos,
                dias_semana=dias if dias != rota.dias_semana else None,
                horarios_por_dia=por_dia, turno_bloco=a.turno_bloco)
            b.atividades += 1
            rota.itens.append(item)
        if not rota.itens:
            previa.avisos.append(f'{nome_rota}: nenhuma atividade aproveitável; não entra.')
            continue
        # O bloco principal (mais minutos) dá o nome à rota.
        por_bloco = Counter()
        for it in rota.itens:
            por_bloco[it.bloco.rotulo] += it.tempo_previsto_min or 0
        principal = max(por_bloco, key=lambda k: (por_bloco[k], -list(por_bloco).index(k)))
        rota.nome = f'{nome_rota} — {principal}'[:150]
        previa.rotas.append(rota)
    previa.blocos = list(blocos.values())
    previa.checklists = list(checklists.values())
    if local is not None:
        _marcar_existentes(previa, local)
    return previa


def _marcar_existentes(previa, local):
    from facilities.cadastros.models import Ambiente, Setor
    ambientes = {(_normal(a.setor.nome), _normal(a.nome)) for a in Ambiente.query.join(Setor)
                 .filter(Setor.client_id == local.id, Ambiente.is_active.is_(True))}
    for b in previa.blocos:
        b.existe = b.chave in ambientes
    nomes = {_normal(m.nome) for m in ChecklistModelo.query.filter_by(client_id=local.id)}
    for c in previa.checklists:
        c.existe = _normal(c.nome) in nomes
    rotas = {_normal(r.nome) for r in Rota.query.filter_by(client_id=local.id)}
    for r in previa.rotas:
        r.existe = _normal(r.nome) in rotas


def _rota_existente(local, nome):
    # Nome exato primeiro: o lower() do SQLite não conhece acento ("GALPÃO").
    return (Rota.query.filter_by(client_id=local.id, nome=nome).first()
            or Rota.query.filter(Rota.client_id == local.id,
                                 db.func.lower(Rota.nome) == nome.lower()).first())


def rotas_que_ja_existem(previa, local):
    return [r.nome for r in previa.rotas if _rota_existente(local, r.nome) is not None]


def descrever(previa):
    """A prévia em texto corrido (o ``--dry-run`` e o log)."""
    linhas = [f'BLOCOS ({len(previa.blocos)}) — um ponto com QR por bloco:']
    for b in previa.blocos:
        linhas.append(f'  - {b.rotulo} [{b.setor}] · {b.atividades} atividade(s)'
                      + (' · já existe' if b.existe else ''))
    linhas.append(f'CHECKLISTS ({len(previa.checklists)}):')
    for c in previa.checklists:
        linhas.append(f'  - {c.nome}' + (' · já existe' if c.existe else ''))
        linhas.append('      perguntas: ' + ' | '.join(i['texto'] for i in c.itens))
        if c.materiais:
            linhas.append(f'      materiais: {c.materiais}')
        if c.epi:
            linhas.append(f'      EPI: {c.epi}')
    linhas.append(f'ROTAS ({len(previa.rotas)}):')
    for r in previa.rotas:
        dias = ', '.join(DIAS_DA_SEMANA[d] for d in r.dias_semana)
        linhas.append(f'  - {r.nome} · {r.turno} {r.hora_inicio}–{r.hora_fim} · {dias}'
                      + (f' · intervalo {r.intervalo_inicio}–{r.intervalo_fim}'
                         if r.intervalo_inicio else '')
                      + (' · JÁ EXISTE' if r.existe else ''))
        if r.responsavel:
            linhas.append(f'      Responsável: {r.responsavel} (escalar depois)')
        for it in r.itens:
            extra = ''
            if it.dias_semana:
                extra += ' · só ' + ', '.join(DIAS_DA_SEMANA[d] for d in it.dias_semana)
            for d, h in it.horarios_por_dia.items():
                extra += f' · {DIAS_DA_SEMANA[d]} às {h}'
            linhas.append(f'      {it.ordem:>2}. {it.horario} · {it.tempo_previsto_min or "—"} '
                          f'min · {it.bloco.rotulo} · {it.checklist.nome}{extra}')
        if r.observacoes:
            linhas.append('      Observações: ' + r.observacoes.replace('\n', ' '))
    if previa.avisos:
        linhas.append(f'AVISOS ({len(previa.avisos)}):')
        linhas += [f'  ! {a}' for a in previa.avisos]
    return '\n'.join(linhas)


# ── 3. Aplicação ──────────────────────────────────────────────────────────

def _checklist_do_local(local, previsto, autor):
    """O checklist do local com este nome (reaproveitado e atualizado se o
    plano mudou — ``publicar`` só cria versão quando algo muda) ou um novo."""
    m = (ChecklistModelo.query.filter_by(client_id=local.id, nome=previsto.nome).first()
         or ChecklistModelo.query.filter(
             ChecklistModelo.client_id == local.id,
             db.func.lower(ChecklistModelo.nome) == previsto.nome.lower()).first())
    novo = m is None
    if novo:
        m = servico.criar_modelo(local.id, {'nome': previsto.nome,
                                            'descricao': previsto.descricao}, autor)
    m.is_active = True
    servico.publicar(m, previsto.itens, {'descricao': previsto.descricao,
                                         'materiais': previsto.materiais, 'epi': previsto.epi},
                     autor)
    return m, novo


def aplicar(previa, local, autor, substituir=False, pular_existentes=False):
    """Grava no local o que a prévia mostrou. → resumo (contagens).

    Rota com o mesmo nome já no local: recusa (nada é gravado), a não ser
    ``substituir`` (a antiga é desativada e renomeada) ou ``pular_existentes``
    (a rota do plano é deixada de lado; o resto entra). O commit é de quem
    chama."""
    from facilities.cadastros.servico import CadastroInvalido, registrar
    from facilities.rotas.rotas_prontas import _ambiente, _ponto, _setor
    resumo = {'setores': 0, 'ambientes': 0, 'pontos': 0, 'checklists': 0,
              'checklists_atualizados': 0, 'rotas': 0, 'itens': 0, 'substituidas': 0,
              'puladas': 0}
    existentes = rotas_que_ja_existem(previa, local)
    if existentes and not substituir and not pular_existentes:
        raise RegraDaRota('Já existe neste local: ' + ', '.join(f'"{n}"' for n in existentes)
                          + '. Para trocar pelo plano novo, marque "Substituir as rotas que já '
                          'existem" (a antiga é desativada; o histórico fica).')
    try:
        pontos = {}
        for b in previa.blocos:
            setor, novo = _setor(local, b.setor, autor)
            resumo['setores'] += novo
            amb, novo = _ambiente(local, setor, b.nome, autor)
            resumo['ambientes'] += novo
            ponto, novo = _ponto(local, amb, autor)
            resumo['pontos'] += novo
            pontos[b.chave] = ponto
        modelos = {}
        for c in previa.checklists:
            m, novo = _checklist_do_local(local, c, autor)
            resumo['checklists' if novo else 'checklists_atualizados'] += 1
            modelos[id(c)] = m
        hoje = datetime.now().strftime('%d/%m/%Y %H:%M')
        for r in previa.rotas:
            antiga = _rota_existente(local, r.nome)
            if antiga is not None:
                if pular_existentes and not substituir:
                    resumo['puladas'] += 1
                    continue
                antiga.is_active = False
                antiga.nome = f'{r.nome[:110]} (substituída em {hoje})'[:150]
                db.session.flush()
                resumo['substituidas'] += 1
            rota = servico.salvar_rota(local, None, {
                'nome': r.nome, 'turno': r.turno, 'hora_inicio': r.hora_inicio,
                'hora_fim': r.hora_fim, 'tolerancia_min': r.tolerancia_min,
                'dias_semana': r.dias_semana, 'intervalo_inicio': r.intervalo_inicio,
                'intervalo_fim': r.intervalo_fim, 'observacoes': r.observacoes,
                'presencas_por_turno': '0'})
            for it in r.itens:
                servico.salvar_item(rota, None, {
                    'ponto_id': pontos[it.bloco.chave].id, 'modelo_id': modelos[id(it.checklist)].id,
                    'horarios': it.horario, 'tempo_previsto_min': it.tempo_previsto_min or '',
                    'dias_semana': it.dias_semana or '', 'horarios_por_dia': it.horarios_por_dia,
                    'ordem': str(it.ordem)})
                resumo['itens'] += 1
            registrar('rota', rota, 'importada', None, autor, client_id=local.id)
            resumo['rotas'] += 1
    except CadastroInvalido as erro:
        raise RegraDaRota(str(erro))
    db.session.flush()
    return resumo


# ── A planilha enviada pela tela ──────────────────────────────────────────

def guardar_planilha(local, arquivo):
    """Confere (xlsx, até 5 MB, assinatura ZIP) e guarda a planilha enviada;
    → (chave, plano lido). A chave volta no "Criar"."""
    from facilities.core.anexos import ANEXO_ASSINATURA, storage
    if arquivo is None or not getattr(arquivo, 'filename', ''):
        raise RegraDaRota('Escolha o arquivo .xlsx do plano de limpeza.')
    nome = arquivo.filename.replace('\\', '/').rsplit('/', 1)[-1].strip()
    ext = nome.rsplit('.', 1)[-1].lower() if '.' in nome else ''
    dados = arquivo.read(PLANILHA_MAX + 1)
    if ext != 'xlsx' or not dados or len(dados) > PLANILHA_MAX \
            or not any(dados.startswith(a) for a in ANEXO_ASSINATURA['xlsx']):
        raise RegraDaRota('Envie o plano em Excel (.xlsx, até 5 MB).')
    plano = ler(dados)
    chave = f'planos/{local.id}/{uuid.uuid4().hex}.xlsx'
    storage().save(chave, dados, 'application/vnd.openxmlformats-officedocument.'
                                 'spreadsheetml.sheet')
    return chave, plano


def ler_guardada(local, chave):
    from facilities.core.anexos import storage
    chave = str(chave or '').strip()
    m = CHAVE_DA_PLANILHA.match(chave)
    if m is None or int(m.group(1)) != local.id:
        raise RegraDaRota('A prévia venceu. Envie a planilha de novo.')
    try:
        dados = storage().load(chave)
    except Exception:
        raise RegraDaRota('A prévia venceu. Envie a planilha de novo.')
    return ler(dados)


# ── Cadastro de teste: Trino Frio ─────────────────────────────────────────

NOME_CLIENTE_TRINO_FRIO = 'Trino Frio'
NOME_LOCAL_TRINO_FRIO = 'Galpão Frio — Trino Frio'


def caminho_da_planilha_trino_frio():
    import os
    import facilities
    return os.path.join(os.path.dirname(os.path.abspath(facilities.__file__)),
                        PLANILHA_TRINO_FRIO)


def semear_trino_frio(autor=None):
    """Empresa (a primeira ativa, ou "Trino Serviços"), cliente "Trino Frio"
    (centro de custo TF), local "Galpão Frio — Trino Frio" e o plano da
    planilha empacotada. Rodar de novo não duplica: rota que já existe é
    pulada. Nome de pessoa não entra. → (local, resumo)."""
    from facilities.cadastros import servico as cad
    from facilities.cadastros.models import Contratante, Empresa
    from facilities.cadastros.servico import registrar_local
    from facilities.core.models import Client
    emp = Empresa.query.filter_by(is_active=True).order_by(Empresa.id).first()
    if emp is None:
        emp = cad.salvar_empresa(None, {'nome': 'Trino Serviços'}, autor)
        db.session.flush()
    cli = Contratante.query.filter(
        db.func.lower(Contratante.nome) == NOME_CLIENTE_TRINO_FRIO.lower()).first()
    if cli is None:
        cli = cad.salvar_contratante(None, {'nome': NOME_CLIENTE_TRINO_FRIO,
                                            'empresa_id': emp.id, 'centro_custo': 'TF'}, autor)
        db.session.flush()
    loc = Client.query.filter(db.func.lower(Client.name) == NOME_LOCAL_TRINO_FRIO.lower(),
                              Client.is_deleted.is_(False)).first()
    if loc is None:
        loc = Client(name=NOME_LOCAL_TRINO_FRIO, cost_center='TF-01', is_active=True,
                     contratante_id=cli.id)
        db.session.add(loc)
        db.session.flush()
        registrar_local(loc, None, autor, nova=True)
    elif loc.contratante_id is None:
        loc.contratante_id = cli.id
    previa = planejar(ler(caminho_da_planilha_trino_frio()), loc)
    resumo = aplicar(previa, loc, autor, pular_existentes=True)
    resumo['avisos'] = previa.avisos
    return loc, resumo
