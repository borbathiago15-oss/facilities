"""Relatório mensal do cliente (decisões do dono, 18/10/2026): tela + PDF +
Excel; fotos em amostra automática que a gestão pode trocar; o cliente só vê
o relatório liberado; nada de custo nem nome da equipe."""
import io
from datetime import date, datetime, time, timedelta

import pytest

from facilities import tempo
from tests import fabrica
from tests.test_rotas_servico import Arquivo, _png

SETEMBRO = date(2026, 9, 1)


def _utc(dia, hhmm):
    h, m = map(int, hhmm.split(':'))
    return tempo.para_utc(dia, time(h, m))


@pytest.fixture()
def cena(app_module):
    from app import User, db
    from facilities.chamados import servico as chs
    from facilities.chamados.models import ChamadoAnexo
    from facilities.core.anexos import storage
    from facilities.rotas import servico as s
    with app_module.app.app_context():
        ids = fabrica.estrutura(app_module)
        loc = fabrica.local()
        asg = fabrica.pessoa(app_module, 'asg.rel@trino.com', 'encarregado',
                             locais=['Shopping Centro'], nome='Ana Limpeza')
        coord = fabrica.pessoa(app_module, 'coord.rel@trino.com', 'pleno',
                               locais=['Shopping Centro'], nome='Carla Coordenadora')
        fabrica.pessoa(app_module, 'sup.rel@trino.com', 'supervisor', locais=['Shopping Centro'])
        fabrica.pessoa(app_module, 'fora.rel@trino.com', 'pleno', locais=['Hospital Norte'])
        m = s.criar_modelo(loc.id, {'nome': 'Banheiro'}, None)
        s.publicar(m, s.normalizar_itens([{'texto': 'Piso limpo?', 'tipo': 'sim_nao',
                                           'problema_se': 'nao'}]), {'exige_foto_depois': '1'},
                   None)
        rota = s.salvar_rota(loc, None, {'nome': 'Limpeza', 'turno': 'T1', 'hora_inicio': '07:00',
                                         'hora_fim': '15:00', 'tolerancia_min': '30',
                                         'dias_semana': '1234567'})
        s.salvar_item(rota, None, {'ponto_id': ids['ponto'], 'modelo_id': m.id,
                                   'horarios': '09:00'})
        db.session.commit()
        for dia in (date(2026, 9, 14), date(2026, 9, 15)):
            s.gerar(dia, agora_utc=_utc(dia, '00:00'))
        db.session.commit()
        from facilities.rotas.models import TarefaDia
        t14 = TarefaDia.query.filter_by(data=date(2026, 9, 14)).one()
        u = db.session.get(User, asg)
        s.registrar_execucao(u, {'envio_id': 'rel-0000001', 'ponto': t14.ponto.codigo_qr,
                                 'tarefa_id': t14.id, 'respostas': {'piso_limpo': 'nao'},
                                 'sem_sinal': True, 'lido_em': _utc(t14.data, '09:05').isoformat()},
                             {'depois': [Arquivo()]}, agora_utc=_utc(t14.data, '09:06'))
        s.fechar_perdidas(_utc(date(2026, 9, 16), '00:00'))
        ch = chs.abrir(db.session.get(User, coord), loc.id, 'Lâmpada queimada')
        db.session.flush()
        ch.aberto_em = _utc(date(2026, 9, 10), '08:00')
        ch.prazo = _utc(date(2026, 9, 13), '08:00')
        ch.status, ch.solucao = 'fechado', 'Trocada'
        ch.resolvido_em = _utc(date(2026, 9, 11), '08:00')
        storage().save('teste/rel-depois.png', _png(), 'image/png')
        db.session.add(ChamadoAnexo(chamado_id=ch.id, nome='d.png', mime='image/png',
                                    storage_key='teste/rel-depois.png', momento='depois'))
        db.session.commit()
        ids.update(asg=asg, coord=coord, chamado=ch.id)
    return ids


def _dados(app_module, cena):
    from app import Client, db
    from facilities.relatorios import cliente as rel
    return rel.dados(db.session.get(Client, cena['local']), SETEMBRO,
                     agora=_utc(date(2026, 10, 18), '12:00'))


def test_numeros_do_mes(app_module, cena):
    with app_module.app.app_context():
        d = _dados(app_module, cena)
        assert d['nome_do_mes'] == 'setembro de 2026'
        rot = d['rotinas'][0]
        # As NCs do checklist de 14/09 abriram chamado (bloco A) — conta como aberto.
        assert (rot['previstas'], rot['no_prazo'], rot['nao_feitas']) == (2, 1, 1)
        assert rot['pct'] == 50.0
        c = d['chamados']
        assert c['resolvidos'] == 1 and c['pct_no_prazo'] == 100.0
        assert c['tempo_medio_h'] == 24.0
        assert d['ncs']['total'] == 1 and d['ncs']['itens'] == [('Piso limpo?', 1)]


def test_fotos_amostra_e_escolha(app_module, cena):
    from app import User, db
    from facilities.relatorios import cliente as rel
    with app_module.app.app_context():
        todas = rel.candidatas(cena['local'], SETEMBRO)
        assert {f['tipo'] for f in todas} == {'execucao', 'chamado'}
        assert all('Ana' not in f['legenda'] for f in todas)
        fotos, automatica = rel.fotos_do_relatorio(None, todas)
        assert automatica and len(fotos) == 2
        coord = db.session.get(User, cena['coord'])
        r = rel.escolher_fotos(cena['local'], SETEMBRO, [todas[0]['ref'], 'execucao:99999'], coord)
        db.session.commit()
        fotos, automatica = rel.fotos_do_relatorio(r, todas)
        assert not automatica and [f['ref'] for f in fotos] == [todas[0]['ref']]
        with pytest.raises(rel.RegraDoRelatorio):
            rel.escolher_fotos(cena['local'], SETEMBRO, [], db.session.get(User, cena['asg']))


def test_liberar_so_mes_fechado_e_so_a_gestao(app_module, cena):
    from app import User, db
    from facilities.relatorios import cliente as rel
    with app_module.app.app_context():
        coord = db.session.get(User, cena['coord'])
        with pytest.raises(rel.RegraDoRelatorio, match='ainda não acabou'):
            rel.liberar(cena['local'], tempo.hoje().replace(day=1), coord)
    sup = fabrica.entrar(app_module, 'sup.rel@trino.com')
    url = f"/relatorios/cliente/{cena['local']}/2026-09"
    assert sup.get(url).status_code == 200
    sup.post(url + '/liberar')
    with app_module.app.app_context():
        assert rel.relatorio_do_mes(cena['local'], SETEMBRO) is None
    fora = fabrica.entrar(app_module, 'fora.rel@trino.com')
    assert fora.get(url).status_code == 403
    assert fora.get(url + '/arquivo.pdf').status_code == 403


def test_ciclo_gestao_e_cliente(app_module, cena):
    from facilities.relatorios import cliente as rel
    coord = fabrica.entrar(app_module, 'coord.rel@trino.com')
    url = f"/relatorios/cliente/{cena['local']}/2026-09"
    pag = coord.get(url).get_data(as_text=True)
    assert 'Rotinas por ambiente' in pag and 'Lâmpada queimada' in pag
    assert 'Liberar para o cliente' in pag
    coord.post(url + '/recado', data={'recado': 'Reforçamos o 2º andar.'})
    # Antes de liberar, o cliente não vê.
    app_module.app.config['CLIENTE_EXTERNO_LIBERADO'] = True
    with app_module.app.app_context():
        fabrica.cliente(app_module, 'cli.rel@gl.com', cena['cliente'])
    cli = fabrica.entrar_com_codigo(app_module, 'cli.rel@gl.com')
    url_cli = f"/cliente/relatorios/{cena['local']}/2026-09"
    assert cli.get(url_cli).status_code == 404
    assert cli.get(url_cli + '/arquivo.pdf').status_code == 404
    coord.post(url + '/liberar')
    with app_module.app.app_context():
        assert rel.relatorio_do_mes(cena['local'], SETEMBRO).liberado_em is not None
    assert 'Setembro de 2026' in cli.get('/cliente/relatorios').get_data(as_text=True)
    pag = cli.get(url_cli).get_data(as_text=True)
    assert 'Reforçamos o 2º andar.' in pag and 'Rotinas por ambiente' in pag
    for proibido in ('Ana Limpeza', 'Carla Coordenadora', 'R$', 'Liberar para o cliente'):
        assert proibido not in pag, proibido
    # PDF e Excel.
    from pypdf import PdfReader
    r = cli.get(url_cli + '/arquivo.pdf')
    assert r.mimetype == 'application/pdf'
    texto = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(r.data)).pages)
    assert 'Relatório mensal' in texto and 'setembro de 2026' in texto
    assert 'Ana Limpeza' not in texto and 'Carla' not in texto
    from openpyxl import load_workbook
    r = cli.get(url_cli + '/arquivo.xlsx')
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[:3] == ['Resumo', 'Rotinas por ambiente', 'Chamados']
    # Recolher tira do portal.
    coord.post(url + '/recolher')
    assert cli.get(url_cli).status_code == 404


def test_mes_invalido_vira_mes_passado():
    from facilities.relatorios.cliente import mes_de
    assert mes_de('abc', hoje=date(2026, 10, 18)) == SETEMBRO
    assert mes_de('2026-07', hoje=date(2026, 10, 18)) == date(2026, 7, 1)
